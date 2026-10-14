"""
Genereer een fictief selectiebestand voor de bachelor Geneeskunde (demo).

Bedoeld om de tool te laten zien aan geneeskundeopleidingen van verschillende
universiteiten. De vorm (twee rondes, collegetoets in ronde 1) komt uit de
R-analyses in de map "Analyse Selectie Geneeskunde" (niet in de repo); de
instrumenten hebben algemene namen, zodat elke opleiding er iets van de eigen
procedure in herkent. De instellingsnaam is fictief.

- Ronde 1 (alle aanmelders): toetsen in % goed. Een collegetoets
  (proefstuderen: studiestof bestuderen, dan een toets), een kennistoets
  bèta, een wetenschappelijk artikel lezen, een capaciteitentest (verbaal en
  numeriek) en een situational judgement test. Samen een Z-score ronde 1;
  alleen de besten gaan door.
- Ronde 2 (alleen wie door ronde 1 kwam): drie MMI-stations (multiple
  mini-interviews), elk beoordeeld op 1-10.
- Eindscore en rangnummer bepalen wie een plaats krijgt.

De uitkomsten uit de R-analyses (gemiddeld cijfer t/m periode 7,
Professioneel Gedrag in de master in een keer gehaald) laten we weg: de tool
kijkt naar doorstroom naar jaar 2.

Twee rondes in een blad: wie in ronde 1 afviel heeft geen ronde-2-scores. De
regressie rekent alleen met gestarte studenten, en die hebben beide rondes.
De Z-score ronde 1 en het MMI-gemiddelde vatten andere kolommen samen en
staan daarom uitgevinkt in de config, net als de subtotalen bij Leiden.

500 aanmelders, 260 door naar ronde 2, 180 plaatsen. Bachelor, header_rij=2.
Niet herleidbaar naar echte data.

Draai:
    uv run python scripts/eenmalig/maak_fictief_demo_geneeskunde.py
"""

import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from cho_transform import bouw_ruwe_cho  # noqa: E402
from config_wizard import exporteer_config_excel  # noqa: E402

RNG = np.random.default_rng(2468)

OUT_DIR = Path("data/fictief")
OUT_DIR.mkdir(parents=True, exist_ok=True)

N = 500
N_RONDE_2 = 260
N_PLAATSEN = 180
OPLEIDING = "Geneeskunde"
INSTELLING = "Universiteit Middelland"  # fictief
JAAR = 2026
BLAD_NAAM = "Selectie Geneeskunde"

RONDE_1 = "Ronde 1: toetsen"
RONDE_2 = "Ronde 2: MMI"


def z(x):
    return (x - np.nanmean(x)) / np.nanstd(x)


def procent(latent, gem, sd):
    return np.clip(np.round(gem + sd * latent), 0, 100)


studentnummers = RNG.choice(range(2600000, 2699999), size=N, replace=False)
studentnummers.sort()

# Drie onderliggende eigenschappen: studievaardigheid (wat de cognitieve
# toetsen meten), persoonlijke vaardigheden (wat de SJT en de MMI meten) en
# studie-inzet (wat vooral proefstuderen zichtbaar maakt, en wat het meest
# bepaalt of iemand doorstroomt).
vaardigheid = RNG.normal(0, 1, N)
persoonlijk = RNG.normal(0, 1, N)
inzet = RNG.normal(0, 1, N)


def meting(lading_v, lading_p=0.0, lading_i=0.0):
    ruis = np.sqrt(1 - lading_v**2 - lading_p**2 - lading_i**2)
    return (
        lading_v * vaardigheid
        + lading_p * persoonlijk
        + lading_i * inzet
        + ruis * RNG.normal(0, 1, N)
    )


# -- Ronde 1: toetsen (% goed) ------------------------------------------------
# kolomnaam: (scores, item, criterium)
ronde1 = {
    "Collegetoets % goed": (
        procent(meting(0.55, 0, 0.6), 64, 13),
        "Collegetoets (proefstuderen)",
        "Studievaardigheid",
    ),
    "Kennistoets bèta % goed": (
        procent(meting(0.65), 58, 14),
        "Kennistoets bèta",
        "Vakkennis",
    ),
    "Wetenschappelijk artikel % goed": (
        procent(meting(0.55), 61, 12),
        "Wetenschappelijk artikel lezen",
        "Academische vaardigheden",
    ),
    "Capaciteitentest verbaal % goed": (
        procent(meting(0.5), 56, 15),
        "Capaciteitentest verbaal",
        "Cognitieve capaciteit",
    ),
    "Capaciteitentest numeriek % goed": (
        procent(meting(0.5), 53, 16),
        "Capaciteitentest numeriek",
        "Cognitieve capaciteit",
    ),
    "SJT % goed": (
        procent(meting(0.15, 0.5), 62, 11),
        "Situational judgement test",
        "Professioneel handelen",
    ),
}

z_ronde1 = z(np.mean([z(v[0]) for v in ronde1.values()], axis=0)).round(2)
door = pd.Series(z_ronde1).rank(ascending=False, method="first").values <= N_RONDE_2

# -- Ronde 2: MMI-stations (1-10), alleen wie door ronde 1 kwam ---------------
ronde2 = {
    "MMI Communicatie (1-10)": ("MMI-station communicatie", "Communicatie"),
    "MMI Ethisch redeneren (1-10)": (
        "MMI-station ethisch redeneren",
        "Ethisch redeneren",
    ),
    "MMI Samenwerken (1-10)": ("MMI-station samenwerken", "Samenwerken"),
}
mmi = {}
for naam in ronde2:
    score = np.clip(np.round((6.6 + 1.2 * meting(0.2, 0.6)) * 2) / 2, 1, 10)
    mmi[naam] = np.where(door, score, np.nan)
mmi_gem = np.full(N, np.nan)
mmi_gem[door] = np.mean([v[door] for v in mmi.values()], axis=0).round(2)

# -- Eindscore, rangnummer en besluit ----------------------------------------
eindscore = np.full(N, np.nan)
eindscore[door] = ((z_ronde1[door] + z(mmi_gem[door])) / 2).round(2)
rang = np.full(N, np.nan)
rang[door] = pd.Series(eindscore[door]).rank(ascending=False, method="first").values
plaats = door & (rang <= N_PLAATSEN)

besluit = np.where(
    plaats,
    "Plaats aangeboden",
    np.where(door, "Afgewezen ronde 2", "Afgewezen ronde 1"),
)

df = pd.DataFrame({"Studentnummer": studentnummers})
for naam, (scores, _, _) in ronde1.items():
    df[naam] = scores
df["Z-score ronde 1"] = z_ronde1
df["Uitslag ronde 1"] = np.where(door, "Door naar ronde 2", "Afgewezen")
for naam, scores in mmi.items():
    df[naam] = scores
df["MMI gemiddeld"] = mmi_gem
df["Eindscore"] = eindscore
df["Rangnummer"] = rang
df["Besluit"] = besluit

print(f"Selectiebestand: {df.shape[0]} aanmelders, {df.shape[1]} kolommen")

# -- Excel met een kopregel per ronde (header_rij=2) --------------------------
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

xlsx_path = OUT_DIR / "selectiedata_demo_geneeskunde_2026.xlsx"

wb = Workbook()
ws = wb.active
ws.title = BLAD_NAAM

header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
group_font = Font(bold=True, size=12)

kolommen = list(df.columns)
koppen = {
    "Collegetoets % goed": "Ronde 1: toetsen (alle aanmelders)",
    "MMI Communicatie (1-10)": "Ronde 2: multiple mini-interviews (MMI)",
    "Eindscore": "Besluit",
}
for kolom, kop in koppen.items():
    ws.cell(row=1, column=kolommen.index(kolom) + 1, value=kop).font = group_font

for ci, naam in enumerate(kolommen, 1):
    cell = ws.cell(row=2, column=ci, value=naam)
    cell.font = header_font
    cell.fill = header_fill

for ri, (_, row) in enumerate(df.iterrows(), 3):
    for ci, val in enumerate(row, 1):
        if pd.notna(val):
            ws.cell(row=ri, column=ci, value=val)

wb.save(xlsx_path)
print(f"Opgeslagen: {xlsx_path}")

# -- Configuratiebestand ------------------------------------------------------
config_kolommen = []
for naam in kolommen:
    if naam in ronde1:
        _, item, criterium = ronde1[naam]
        rij = {"instrument": RONDE_1, "item": item, "criterium": criterium}
        config_kolommen.append(
            {"meenemen": True, "kolom_naam": naam, "schaal": "0-100", **rij}
        )
    elif naam in ronde2:
        item, criterium = ronde2[naam]
        rij = {"instrument": RONDE_2, "item": item, "criterium": criterium}
        config_kolommen.append(
            {"meenemen": True, "kolom_naam": naam, "schaal": "1-10", **rij}
        )
    elif naam in ("Z-score ronde 1", "MMI gemiddeld"):
        # Vatten de losse toetsen/stations samen: zouden ze dubbel tellen.
        instrument = RONDE_1 if naam == "Z-score ronde 1" else RONDE_2
        config_kolommen.append(
            {
                "meenemen": False,
                "kolom_naam": naam,
                "instrument": instrument,
                "item": naam,
            }
        )
    else:
        config_kolommen.append({"meenemen": False, "kolom_naam": naam})

config_path = OUT_DIR / "config_demo_geneeskunde_2026.xlsx"
config_path.write_bytes(
    exporteer_config_excel(
        {
            "koppel_id_kolom": "Studentnummer",
            "opleiding": OPLEIDING,
            "instellingscode": INSTELLING,
            "jaar": str(JAAR),
            "blad_naam": BLAD_NAAM,
            "header_rij": "2",
            "totaalscore_kolom": "Eindscore",
            "kolommen": config_kolommen,
        }
    )
)

# -- 1CijferHO-data ----------------------------------------------------------------
# Niet iedereen met een plaats schrijft zich in.
ingeschreven = plaats & (RNG.random(N) < 0.92)
ids = studentnummers[ingeschreven]
n_ingeschreven = len(ids)

# Geneeskunde houdt veel studenten vast; wie hoger scoorde, iets vaker.
lin = (
    2.0
    + 0.35 * z(vaardigheid[ingeschreven])
    + 0.25 * z(persoonlijk[ingeschreven])
    + 0.6 * z(inzet[ingeschreven])
)
doorstroomt = RNG.random(n_ingeschreven) < 1 / (1 + np.exp(-lin))

geslacht = RNG.choice(["vrouw", "man"], size=n_ingeschreven, p=[0.70, 0.30])
vooropleiding = RNG.choice(
    ["VWO", "WO bachelor", "HBO propedeuse"],
    size=n_ingeschreven,
    p=[0.86, 0.08, 0.06],
)

cho_df = bouw_ruwe_cho(
    ids,
    jaar=JAAR,
    doorstroomt=doorstroomt,
    opleiding=OPLEIDING,
    instellingscode=INSTELLING,
    geslacht=geslacht,
    vooropleiding_omschrijving=vooropleiding,
)
cho_path = OUT_DIR / "1cho_data_demo_geneeskunde_2026.csv"
cho_df.to_csv(cho_path, index=False, sep=";")

# -- Kopieer naar demo --------------------------------------------------------
demo_subdir = Path("data/demo/demo_geneeskunde_2026")
demo_subdir.mkdir(parents=True, exist_ok=True)
shutil.copy2(xlsx_path, demo_subdir / "selectiedata.xlsx")
shutil.copy2(config_path, demo_subdir / "config.xlsx")
shutil.copy2(cho_path, demo_subdir / "1cho_data.csv")
print(f"Demo bestanden in {demo_subdir}/")

print("\nSamenvatting:")
print(f"  Aanmelders: {N}, door naar ronde 2: {int(door.sum())}")
print(f"  Plaats aangeboden: {int(plaats.sum())}, ingeschreven: {n_ingeschreven}")
print(f"  Doorgestroomd: {int(doorstroomt.sum())}")
print(f"  Niet doorgestroomd: {int(n_ingeschreven - doorstroomt.sum())}")
print(
    f"  Config: {len(ronde1) + len(ronde2)} items "
    f"({len(ronde1)} ronde 1, {len(ronde2)} ronde 2), "
    "Z-score ronde 1 en MMI gemiddeld uitgevinkt"
)
